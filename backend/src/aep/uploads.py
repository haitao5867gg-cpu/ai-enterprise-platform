"""Validate a batch of source files and stage its Parquet tables atomically."""

import csv
import io
import json
import re
import shutil
import tempfile
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

import duckdb
from fastapi import UploadFile
from openpyxl import load_workbook

from aep.config import Settings

MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_FILES = 20
csv.field_size_limit(MAX_FILE_BYTES)


class UploadError(ValueError):
    """A validation error safe to show to the uploader."""


def table_name(original: str, taken: set[str]) -> str:
    base = re.sub(r"[^a-z0-9]+", "_", original.lower()).strip("_")[:50] or "table"
    if base[0].isdigit():
        base = "table_" + base
    candidate, number = base, 1
    while candidate in taken:
        number += 1
        candidate = f"{base}_{number}"
    taken.add(candidate)
    return candidate


def _headers(values: tuple[Any, ...] | list[Any]) -> list[str]:
    if not values or any(not isinstance(v, str) or not v.strip() for v in values):
        raise UploadError("第一行必须是完整、非空的表头")
    names = [v.strip() for v in values]
    if all(re.fullmatch(r"[-+]?\d+(\.\d+)?", name) for name in names):
        raise UploadError("缺少表头，第一行不能全是数字")
    # DuckDB identifiers are case insensitive, including quoted identifiers.
    if len({name.casefold() for name in names}) != len(names):
        raise UploadError("表头重复（忽略大小写和首尾空格）")
    if any("\x00" in name for name in names):
        raise UploadError("表头不能包含空字符")
    return names


def _csv_rows(content: bytes) -> Iterable[list[str]]:
    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            decoded = content.decode("gbk")
        except UnicodeDecodeError as exc:
            raise UploadError("CSV 编码必须是 UTF-8 或 GBK") from exc
    if "\x00" in decoded:
        raise UploadError("CSV 含有无效的空字符")
    return csv.reader(io.StringIO(decoded, newline=""), strict=True)


def _convert_table(
    rows: Iterable[Iterable[Any]], staging: Path, name: str, conn: duckdb.DuckDBPyConnection
) -> tuple[list[dict[str, Any]], int]:
    iterator = iter(rows)
    header = _headers(list(next(iterator, [])))
    csv_path = staging / f"{name}.csv"
    count = 0
    with csv_path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(header)
        for row in iterator:
            values = list(row)
            if not values or all(v is None or v == "" for v in values):
                continue
            if len(values) != len(header):
                raise UploadError(f"第 {count + 2} 行的列数与表头不一致")
            writer.writerow(values)
            count += 1
    if not count:
        raise UploadError("文件或工作表为空，表头后至少需要一行数据")
    # The CSV is validated and rewritten as UTF-8 before DuckDB infers types.
    relation = conn.read_csv(str(csv_path), header=True, sample_size=-1, delimiter=",")
    columns = [
        {"name": column, "data_type": str(dtype), "nullable": True, "primary_key": False}
        for column, dtype in zip(relation.columns, relation.types, strict=True)
    ]
    relation.write_parquet(str(staging / f"{name}.parquet"), compression="zstd")
    csv_path.unlink()
    return columns, count


@dataclass(frozen=True)
class StagedDataset:
    path: Path
    dataset_id: str
    schema: str
    tables: list[dict[str, Any]]


def stage_dataset(files: list[UploadFile], settings: Settings) -> StagedDataset:
    """No published directory exists until every file and worksheet has succeeded."""
    root = settings.uploads_dir
    root.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=".upload-", dir=root))
    dataset_id = str(uuid4())
    schema = "upload_" + dataset_id[:8]
    tables: list[dict[str, Any]] = []
    taken: set[str] = set()
    filename = "未知文件"
    try:
        with duckdb.connect(config={"threads": "1"}) as conn:
            for upload in files:
                filename = upload.filename or "未命名文件"
                suffix = Path(filename).suffix.lower()
                if suffix not in {".csv", ".xlsx"}:
                    raise UploadError("只支持 .csv 和 .xlsx 文件")
                if upload.size is not None and upload.size > MAX_FILE_BYTES:
                    raise UploadError("单文件不能超过 50 MB")
                content = upload.file.read(MAX_FILE_BYTES + 1)
                if len(content) > MAX_FILE_BYTES:
                    raise UploadError("单文件不能超过 50 MB")
                if not content:
                    raise UploadError("文件为空")

                def add(
                    original: str, rows: Iterable[Iterable[Any]], original_filename: str = filename
                ) -> None:
                    name = table_name(original, taken)
                    columns, count = _convert_table(rows, staging, name, conn)
                    tables.append(
                        {
                            "original_filename": original_filename,
                            "original_name": original,
                            "table": name,
                            "file": f"{name}.parquet",
                            "columns": columns,
                            "row_count": count,
                        }
                    )

                if suffix == ".csv":
                    add(Path(filename).stem, _csv_rows(content))
                else:
                    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
                    try:
                        for sheet in workbook.worksheets:
                            try:
                                add(sheet.title, sheet.iter_rows(values_only=True))
                            except UploadError as exc:
                                raise UploadError(f"工作表「{sheet.title}」：{exc}") from exc
                    finally:
                        workbook.close()
        manifest = {
            "version": 1,
            "dataset_id": dataset_id,
            "schema": schema,
            "uploaded_at": datetime.now(UTC).isoformat(),
            "tables": tables,
        }
        (staging / "manifest.json").write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        # Account for a restrictive API umask: the connector is another, unprivileged user.
        for path in staging.iterdir():
            path.chmod(0o644)
        staging.chmod(0o755)
        return StagedDataset(staging, dataset_id, schema, tables)
    except Exception as exc:
        shutil.rmtree(staging)
        detail = str(exc) if isinstance(exc, UploadError) else "文件解析或转换失败，请检查文件内容"
        raise UploadError(f"文件「{filename}」：{detail}") from exc
