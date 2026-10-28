from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from aep.knowledge.extract import ExtractError, Paragraph, extract


def write_pdf(path: Path, texts: list[str | None]) -> None:
    """Build a tiny text PDF using only pypdf and a standard built-in font."""
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    for text in texts:
        page = writer.add_blank_page(width=300, height=300)
        if text is not None:
            page[NameObject("/Resources")] = DictionaryObject(
                {
                    NameObject("/Font"): DictionaryObject(
                        {NameObject("/F1"): writer._add_object(font)}
                    )
                }
            )
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 20 260 Td ({text}) Tj ET".encode("ascii"))
            page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)


def test_txt_falls_back_to_gb18030(tmp_path):
    path = tmp_path / "资料.txt"
    path.write_bytes("退货期限为三十天。\n\n第二段包括 GB18030 字符𠀀。".encode("gb18030"))
    assert extract(path, "txt") == [
        Paragraph("退货期限为三十天。"),
        Paragraph("第二段包括 GB18030 字符𠀀。"),
    ]


def test_markdown_heading_paths_and_bom(tmp_path):
    path = tmp_path / "资料.md"
    path.write_text(
        "# 退货\n正文\n\n## 期限\n三十天\n### 例外\n说明\n## 方式\n邮寄\n# 换货\n七天",
        encoding="utf-8-sig",
    )
    result = extract(path, "md")
    assert [(p.text, p.heading, p.page) for p in result] == [
        ("正文", "退货", None),
        ("三十天", "退货 > 期限", None),
        ("说明", "退货 > 期限 > 例外", None),
        ("邮寄", "退货 > 方式", None),
        ("七天", "换货", None),
    ]


def test_word_heading_styles_and_tables_remain_in_document_order(tmp_path):
    path = tmp_path / "资料.docx"
    document = Document()
    document.add_heading("退货", level=1)
    document.add_paragraph("第一段")
    document.add_heading("期限", level=2)
    document.add_heading("普通商品", level=3)
    document.add_paragraph("三十天")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "商品", "期限"
    table.cell(1, 0).text, table.cell(1, 1).text = "生鲜", "七天"
    document.styles.add_style("标题 1", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("换货", style="标题 1")
    document.add_paragraph("最后一段")
    document.save(path)
    result = extract(path, "docx")
    assert [p.text for p in result] == [
        "第一段",
        "三十天",
        "商品 | 期限",
        "生鲜 | 七天",
        "最后一段",
    ]
    assert [p.heading for p in result] == [
        "退货",
        "退货 > 期限 > 普通商品",
        "退货 > 期限 > 普通商品",
        "退货 > 期限 > 普通商品",
        "换货",
    ]


def test_pdf_extracts_each_page(tmp_path):
    path = tmp_path / "资料.pdf"
    write_pdf(path, ["Returns within thirty days.", "Contact support to request a refund."])
    assert extract(path, "pdf") == [
        Paragraph("Returns within thirty days.", page=1),
        Paragraph("Contact support to request a refund.", page=2),
    ]


@pytest.mark.parametrize("kind", ["txt", "md", "pdf", "docx"])
def test_empty_and_scanned_documents_report_no_text(tmp_path, kind):
    path = tmp_path / f"空文件.{kind}"
    if kind == "pdf":
        write_pdf(path, [None])
    elif kind == "docx":
        document = Document()
        document.add_table(rows=2, cols=2)
        document.save(path)
    else:
        path.write_text("  \n\t", encoding="utf-8")
    with pytest.raises(ExtractError, match="没有可提取的文字"):
        extract(path, kind)
